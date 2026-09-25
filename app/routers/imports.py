import csv
import io
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, Form, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.core.exceptions import BadRequestError
from app.dependencies.auth import get_current_user
from app.dependencies.database import get_db
from app.models.user import User
from app.schemas.response import ResponseSchema
from app.schemas.word import WordSetCreate
from app.services.word_service import create_word_set

router = APIRouter(prefix="/word-sets/import", tags=["word-sets"])
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500
REQUIRED_HEADERS = {"word", "definition", "example"}


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _read_rows(filename: str, content: bytes) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    extension = Path(filename or "").suffix.casefold()
    if extension not in {".xlsx", ".csv"}:
        raise BadRequestError("Chỉ chấp nhận tệp .xlsx hoặc .csv")
    if len(content) > MAX_FILE_BYTES:
        raise BadRequestError("Tệp vượt quá giới hạn 5 MB")

    try:
        if extension == ".csv":
            text = content.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            values = list(csv.reader(io.StringIO(text), dialect))
        else:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            values = list(workbook.active.iter_rows(values_only=True))
            workbook.close()
    except (UnicodeDecodeError, csv.Error, ValueError, OSError) as exc:
        raise BadRequestError("Không thể đọc tệp. Hãy kiểm tra lại định dạng.") from exc
    if not values:
        raise BadRequestError("Tệp không có dòng tiêu đề")

    headers = [_cell_text(value).casefold().replace(" ", "_") for value in values[0]]
    if "word" not in headers and "term" in headers:
        headers[headers.index("term")] = "word"
    missing_headers = REQUIRED_HEADERS.difference(headers)
    if missing_headers:
        raise BadRequestError(
            "Thiếu cột bắt buộc: " + ", ".join(sorted(missing_headers)),
            details={"required_columns": ["word", "definition", "example"], "found_columns": headers},
        )
    index = {header: headers.index(header) for header in ("word", "definition", "example", "context_sentence") if header in headers}
    valid: list[dict[str, str]] = []
    errors: list[dict[str, Any]] = []
    max_lengths = {"word": 100, "definition": 255, "example": 500, "context_sentence": 1000}
    nonempty_count = 0
    for line_number, row in enumerate(values[1:], start=2):
        if not any(_cell_text(cell) for cell in row):
            continue
        nonempty_count += 1
        if nonempty_count > MAX_ROWS:
            raise BadRequestError(f"Tệp chỉ được có tối đa {MAX_ROWS} dòng dữ liệu")
        data = {key: _cell_text(row[column]) if column < len(row) else "" for key, column in index.items()}
        if not data.get("word") or not data.get("definition"):
            missing = [label for label in ("word", "definition") if not data.get(label)]
            errors.append({"row": line_number, "message": "Thiếu " + " và ".join(missing)})
            continue
        too_long = [
            f"{field} vượt quá {limit} ký tự"
            for field, limit in max_lengths.items()
            if len(data.get(field, "")) > limit
        ]
        if too_long:
            errors.append({"row": line_number, "message": "; ".join(too_long)})
            continue
        valid.append(data)
    return valid, errors


async def _upload_content(file: UploadFile) -> tuple[str, bytes]:
    content = await file.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise BadRequestError("Tệp vượt quá giới hạn 5 MB")
    return file.filename or "", content


@router.post("/preview", response_model=ResponseSchema[dict])
async def preview_word_set_import(
    file: UploadFile = File(...),
    _: User = Depends(get_current_user),
) -> ResponseSchema[dict]:
    filename, content = await _upload_content(file)
    rows, errors = _read_rows(filename, content)
    return ResponseSchema(
        data={"rows": rows, "valid_count": len(rows), "errors": errors},
        message="Đã đọc tệp. Các dòng lỗi sẽ được bỏ qua khi lưu.",
    )


@router.post("", response_model=ResponseSchema[dict], status_code=201)
async def import_word_set(
    file: UploadFile = File(...),
    title: str = Form(min_length=1, max_length=100),
    description: str | None = Form(default=None, max_length=255),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> ResponseSchema[dict]:
    filename, content = await _upload_content(file)
    rows, errors = _read_rows(filename, content)
    if not rows:
        raise BadRequestError("Không có dòng hợp lệ để tạo bộ từ", details={"errors": errors})
    try:
        request = WordSetCreate.model_validate({"title": title.strip(), "description": description, "words": rows})
    except ValidationError as exc:
        raise BadRequestError("Dữ liệu trong tệp vượt quá giới hạn trường", details=exc.errors()) from exc
    word_set = create_word_set(db, request, current_user.id)
    return ResponseSchema(
        data={"word_set": word_set, "imported_count": len(rows), "errors": errors},
        message=f"Đã tạo bộ từ với {len(rows)} từ hợp lệ",
    )


@router.get("/template")
def download_import_template(_: User = Depends(get_current_user)) -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Words"
    sheet.append(["word", "definition", "example"])
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 40
    sheet.column_dimensions["C"].width = 64
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="wordduel-template.xlsx"'},
    )
